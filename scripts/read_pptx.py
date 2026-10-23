import io
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from kb.config import load_nextcloud_settings
from kb.nextcloud_client import NextcloudClient

client = NextcloudClient(load_nextcloud_settings())
prs = Presentation(io.BytesIO(client.download(sys.argv[1])))

for n, slide in enumerate(prs.slides, 1):
    title = slide.shapes.title.text if slide.shapes.title else "(no title)"
    kinds = {"table": 0, "chart": 0, "picture": 0}
    body = []
    for sh in slide.shapes:
        if sh.has_table:
            kinds["table"] += 1
            body.append("TABLE " + str([c.text for c in sh.table.rows[0].cells]))
        elif sh.has_chart:
            kinds["chart"] += 1
            plot = sh.chart.plots[0]
            body.append(f"CHART {list(plot.categories)} {list(plot.series[0].values)}")
        elif sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            kinds["picture"] += 1
        elif sh.has_text_frame and sh != slide.shapes.title:
            body.append(sh.text_frame.text.replace("\n", " / "))
    notes = slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip()
    print(f"{n:2d} {title!r} {kinds} notes={'yes' if notes else 'no'}")
    print("     ", " | ".join(body)[:110])