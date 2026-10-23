import io
import sys

import pandas as pd
from openpyxl import load_workbook

from kb.config import load_nextcloud_settings
from kb.nextcloud_client import NextcloudClient

client = NextcloudClient(load_nextcloud_settings())
data = client.download(sys.argv[1])

wb = load_workbook(io.BytesIO(data), data_only=True)
for ws in wb.worksheets:
    print(f"[{ws.title}] range={ws.dimensions} merged={[str(r) for r in ws.merged_cells.ranges]}")
    for row in ws.iter_rows(min_row=1, max_row=3, values_only=True):
        print("    ", row)

df = pd.read_excel(io.BytesIO(data), sheet_name="Revenue")
print("Revenue_USD total:", df["Revenue_USD"].sum())