import io
import sys

from docx import Document

from kb.config import load_nextcloud_settings
from kb.nextcloud_client import NextcloudClient

client = NextcloudClient(load_nextcloud_settings())
doc = Document(io.BytesIO(client.download(sys.argv[1])))
text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
print("words:", len(text.split()), "| tables:", len(doc.tables), "| images:", len(doc.inline_shapes))
print(text[:300])
for t in doc.tables:
    print("table first row:", [c.text for c in t.rows[0].cells])